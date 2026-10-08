"""Excel (MIS) reports."""
import io
from collections import defaultdict
from decimal import Decimal

from django.db.models import Prefetch
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import access, services
from .models import (
    Agreement, Asset, Compensation, CompensationStatus, CropAssessment, Farmer, LandOwnership,
    LandParcel, Payment,
)

HEADER_FILL = PatternFill("solid", fgColor="1F3A2E")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
MONEY = '#,##,##0.00'
THIN = Side(style="thin", color="D0D5D2")


def _num(v):
    if isinstance(v, Decimal):
        return float(v)
    return v


class Sheet:
    def __init__(self, wb, title, report_title, subtitle, headers, money_cols=(), first=False):
        self.ws = wb.active if first else wb.create_sheet()
        self.ws.title = title[:31]
        self.ws["A1"] = report_title
        self.ws["A1"].font = TITLE_FONT
        self.ws["A2"] = subtitle
        self.ws["A2"].font = Font(italic=True, color="555555")
        self.headers = headers
        self.money_cols = set(money_cols)
        self.row = 4
        for i, h in enumerate(headers, 1):
            c = self.ws.cell(row=self.row, column=i, value=h)
            c.fill, c.font = HEADER_FILL, HEADER_FONT
            c.alignment = Alignment(wrap_text=True, vertical="center")
        self.ws.freeze_panes = self.ws.cell(row=self.row + 1, column=1)
        self.widths = [len(str(h)) + 2 for h in headers]

    def add(self, values, bold=False):
        self.row += 1
        for i, v in enumerate(values, 1):
            c = self.ws.cell(row=self.row, column=i, value=_num(v))
            c.border = Border(bottom=THIN)
            if self.headers[i - 1] in self.money_cols:
                c.number_format = MONEY
            if bold:
                c.font = Font(bold=True)
            self.widths[i - 1] = min(50, max(self.widths[i - 1], len(str(v if v is not None else "")) + 2))

    def finish(self):
        for i, w in enumerate(self.widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = max(10, w)
        if self.row > 4:
            self.ws.auto_filter.ref = f"A4:{get_column_letter(len(self.headers))}{self.row}"


def _subtitle(user, projects):
    names = ", ".join(p.code for p in projects) or "All projects"
    return f"{names}  |  Generated {timezone.localtime():%d-%m-%Y %H:%M} by {user.display_name}"


def _projects(user, project_id):
    qs = access.visible_projects(user)
    return qs.filter(id=project_id) if project_id else qs


def farmers_report(wb, user, projects):
    farmers = access.scope_queryset(user, Farmer.objects.filter(projects__in=projects)).distinct().prefetch_related(
        "projects", Prefetch("ownerships", queryset=LandOwnership.objects.select_related("land")))
    sh = Sheet(wb, "Farmers", "Farmer Master List", _subtitle(user, projects),
               ["Farmer ID", "Name", "Relation", "Mobile", "Village", "Hobli", "Taluk", "District", "Status",
                "KYC", "Aadhaar (last 4)", "Survey numbers", "Projects"], first=True)
    for f in farmers.order_by("village", "farmer_code"):
        sh.add([f.farmer_code, f.name, f"{f.relation_type} {f.relation_name}".strip(), f.mobile, f.village, f.hobli,
                f.taluk, f.district, f.get_status_display(), f.get_kyc_status_display(),
                ("XXXX-XXXX-" + f.aadhaar_last4) if f.aadhaar_last4 else "",
                ", ".join(o.land.survey_label for o in f.ownerships.all()),
                ", ".join(p.code for p in f.projects.all())])
    sh.finish()


def progress_report(wb, user, projects):
    stages = services.active_stages()
    assets = services.assets_with_progress_prefetch(
        Asset.objects.filter(project__in=projects)).order_by("project__code", "line_name", "asset_number")
    sh = Sheet(wb, "Location progress", "Pole / Tower-wise Progress & GPS Report", _subtitle(user, projects),
               ["Project", "Line / Corridor", "Type", "Asset No.", "Survey No.", "Village", "Farmers", "Latitude",
                "Longitude", "GPS accuracy (m)"] + [st.name for st in stages] + ["% complete"], first=True)
    for a in assets:
        prog = services.asset_progress(a, stages)
        lp = a.land_parcel
        sh.add([a.project.code, a.line_name, a.get_asset_type_display(), a.asset_number,
                lp.survey_label if lp else "", lp.village if lp else "",
                ", ".join(f"{f.name} ({f.farmer_code})" for f in services.asset_farmers(a)),
                a.latitude, a.longitude, a.gps_accuracy_m]
               + [("Done" + (f" {i['completed_on']}" if i["completed_on"] else "")) if i["completed"] else (i["detail"] if i["partial"] else "Pending") for i in prog["stages"]]
               + [prog["percent"]])
    sh.finish()


def compensation_report(wb, user, projects):
    comps = services.annotate_paid(Compensation.objects.filter(project__in=projects)).select_related(
        "project", "payee", "asset", "land_parcel", "approved_by").order_by("project__code", "payee__farmer_code")
    money = ["Approved (Rs)", "Paid (Rs)", "Balance (Rs)"]
    sh = Sheet(wb, "Compensation", "Compensation Statement", _subtitle(user, projects),
               ["Project", "Farmer ID", "Payee", "Village", "Survey No.", "Asset No.", "Category", "Description",
                "Status", "Approved on"] + money + ["Payment status"], money_cols=money, first=True)
    tot = [Decimal(0)] * 3
    for c in comps:
        approved = c.approved_amount if c.status == CompensationStatus.APPROVED else Decimal(0)
        paid = c.paid_amount
        vals = [approved, paid, approved - paid]
        tot = [t + v for t, v in zip(tot, vals)]
        sh.add([c.project.code, c.payee.farmer_code, c.payee.name, c.payee.village,
                c.land_parcel.survey_label if c.land_parcel else "", c.asset.asset_number if c.asset else "",
                c.get_category_display(), c.description, c.get_status_display(),
                c.approved_on.strftime("%d-%m-%Y") if c.approved_on else ""] + vals + [c.payment_status.title()])
    sh.add(["TOTAL"] + [""] * 9 + tot + [""], bold=True)
    sh.finish()


def payments_report(wb, user, projects):
    pays = Payment.objects.filter(compensation__project__in=projects).select_related(
        "compensation__project", "compensation__payee", "compensation__asset", "paid_to", "created_by").order_by("payment_date")
    sh = Sheet(wb, "Payments", "Payment Register (recorded payments)", _subtitle(user, projects),
               ["Date", "Project", "Farmer ID", "Payee", "Paid to", "Category", "Asset No.", "Amount (Rs)", "Mode",
                "UTR / Ref No.", "Receipt No.", "Recorded by"], money_cols=["Amount (Rs)"], first=True)
    total = Decimal(0)
    for p in pays:
        total += p.amount
        c = p.compensation
        sh.add([p.payment_date.strftime("%d-%m-%Y"), c.project.code, c.payee.farmer_code, c.payee.name,
                p.paid_to.name if p.paid_to else c.payee.name, c.get_category_display(),
                c.asset.asset_number if c.asset else "", p.amount, p.get_mode_display(), p.reference_number,
                p.receipt_number, p.created_by.display_name if p.created_by else ""])
    sh.add(["TOTAL", "", "", "", "", "", "", total], bold=True)
    sh.finish()


def agreements_report(wb, user, projects):
    ags = Agreement.objects.filter(project__in=projects).select_related("project", "land_parcel").prefetch_related(
        "farmers", "assets").order_by("status", "agreement_number")
    sh = Sheet(wb, "Agreements", "Agreement Register (executed & pending)", _subtitle(user, projects),
               ["Agreement No.", "Project", "Type", "Status", "Farmers", "Survey No.", "Assets", "Agreement date",
                "Period (months)", "Renewal date", "Total consideration (Rs)", "Stamp duty (Rs)", "Registration"],
               money_cols=["Total consideration (Rs)", "Stamp duty (Rs)"], first=True)
    for a in ags:
        sh.add([a.agreement_number, a.project.code, a.get_agreement_type_display(), a.get_status_display(),
                ", ".join(f"{f.name} ({f.farmer_code})" for f in a.farmers.all()),
                a.land_parcel.survey_label if a.land_parcel else "",
                ", ".join(x.asset_number for x in a.assets.all()),
                a.agreement_date.strftime("%d-%m-%Y") if a.agreement_date else "", a.period_months,
                a.renewal_date.strftime("%d-%m-%Y") if a.renewal_date else "", a.total_consideration,
                a.stamp_duty, a.registration_details])
    sh.finish()


def crop_report(wb, user, projects):
    crops = CropAssessment.objects.filter(project__in=projects).select_related("project", "farmer", "land_parcel", "asset")
    sh = Sheet(wb, "Crop assessments", "Crop Compensation Assessments", _subtitle(user, projects),
               ["Date", "Project", "Farmer ID", "Farmer", "Survey No.", "Asset No.", "Season", "Crop", "Area (acres)",
                "Stage", "Revenue assessment (Rs)", "Company assessment (Rs)", "Notes"],
               money_cols=["Revenue assessment (Rs)", "Company assessment (Rs)"], first=True)
    for c in crops:
        sh.add([c.assessment_date.strftime("%d-%m-%Y") if c.assessment_date else "", c.project.code,
                c.farmer.farmer_code, c.farmer.name, c.land_parcel.survey_label if c.land_parcel else "",
                c.asset.asset_number if c.asset else "", c.get_season_display(), c.crop_type, c.crop_area_acres,
                c.get_crop_stage_display(), c.revenue_assessment, c.company_assessment, c.field_inspection_notes])
    sh.finish()


def village_report(wb, user, projects):
    stages = services.active_stages()
    rows = defaultdict(lambda: {"farmers": set(), "lands": set(), "assets": 0, "cleared": 0, "approved": Decimal(0), "paid": Decimal(0), "key": None})
    for a in services.assets_with_progress_prefetch(Asset.objects.filter(project__in=projects)):
        lp = a.land_parcel
        key = (lp.taluk, lp.hobli, lp.village) if lp else ("", "", "(no land linked)")
        r = rows[key]
        r["assets"] += 1
        prog = services.asset_progress(a, stages)
        done = {i["code"] for i in prog["stages"] if i["completed"]}
        r["cleared"] += {"AGREEMENT", "PAYMENT"} <= done
        if lp:
            r["lands"].add(lp.id)
            for f in services.asset_farmers(a):
                r["farmers"].add(f.id)
    for c in services.annotate_paid(Compensation.objects.filter(project__in=projects, status=CompensationStatus.APPROVED)).select_related("payee", "land_parcel"):
        lp = c.land_parcel
        key = (lp.taluk, lp.hobli, lp.village) if lp else (c.payee.taluk, c.payee.hobli, c.payee.village)
        rows[key]["approved"] += c.approved_amount
        rows[key]["paid"] += c.paid_amount
    money = ["Approved (Rs)", "Paid (Rs)", "Balance (Rs)"]
    sh = Sheet(wb, "Village summary", "Village / Hobli / Taluk-wise Summary", _subtitle(user, projects),
               ["Taluk", "Hobli", "Village", "Farmers", "Survey Nos.", "Locations", "ROW cleared"] + money,
               money_cols=money, first=True)
    for key in sorted(rows):
        r = rows[key]
        sh.add(list(key) + [len(r["farmers"]), len(r["lands"]), r["assets"], r["cleared"], r["approved"], r["paid"], r["approved"] - r["paid"]])
    sh.finish()


def farmer_statement(wb, user, farmer):
    projects = access.visible_projects(user)
    ws_sheet = Sheet(wb, "Statement", f"Farmer Statement - {farmer.farmer_code} {farmer.name}",
                     f"{farmer.get_relation_type_display() if farmer.relation_type else ''} {farmer.relation_name}  |  {farmer.village}, {farmer.taluk}, {farmer.district}  |  Mobile {farmer.mobile}  |  KYC: {farmer.get_kyc_status_display()}",
                     ["Date", "Project", "Category", "Particulars", "Approved (Rs)", "Paid (Rs)", "Reference"],
                     money_cols=["Approved (Rs)", "Paid (Rs)"], first=True)
    comps = services.annotate_paid(farmer.compensations.filter(project__in=projects)).select_related("project", "asset", "land_parcel").prefetch_related("payments")
    ta = tp = Decimal(0)
    for c in comps.order_by("approved_on", "created_at"):
        approved = c.approved_amount if c.status == CompensationStatus.APPROVED else Decimal(0)
        ta += approved
        ws_sheet.add([(c.approved_on or c.created_at.date()).strftime("%d-%m-%Y"), c.project.code, c.get_category_display(),
                      " ".join(x for x in [c.description, f"Asset {c.asset.asset_number}" if c.asset else "",
                                            f"Sy.No. {c.land_parcel.survey_label}" if c.land_parcel else "",
                                            f"[{c.get_status_display()}]"] if x), approved, None, ""])
        for p in c.payments.all().order_by("payment_date"):
            tp += p.amount
            ws_sheet.add([p.payment_date.strftime("%d-%m-%Y"), c.project.code, "Payment", f"{p.get_mode_display()} towards {c.get_category_display().lower()}",
                          None, p.amount, p.reference_number])
    ws_sheet.add(["", "", "", "TOTAL", ta, tp, ""], bold=True)
    ws_sheet.add(["", "", "", "BALANCE", ta - tp, None, ""], bold=True)
    ws_sheet.finish()

    lands = LandParcel.objects.filter(ownerships__farmer=farmer, ownerships__is_deleted=False).distinct()
    sh2 = Sheet(wb, "Land & locations", "Land and affected locations", "", ["Survey No.", "Village", "Extent (acres)", "Location", "Project", "Progress"])
    stages = services.active_stages()
    for lp in lands:
        assets = list(services.assets_with_progress_prefetch(lp.assets.filter(project__in=projects)))
        if not assets:
            sh2.add([lp.survey_label, lp.village, lp.total_acres, "", "", ""])
        for a in assets:
            prog = services.asset_progress(a, stages)
            sh2.add([lp.survey_label, lp.village, lp.total_acres, f"{a.get_asset_type_display()} {a.asset_number}", a.project.code,
                     ", ".join(i["name"] for i in prog["stages"] if i["completed"]) or "Not started"])
    sh2.finish()


REPORTS = {
    "farmers": ("Farmer master list", farmers_report),
    "progress": ("Pole/Tower-wise progress & GPS", progress_report),
    "compensation": ("Compensation statement", compensation_report),
    "payments": ("Payment register", payments_report),
    "agreements": ("Agreements (executed & pending)", agreements_report),
    "crop": ("Crop assessments", crop_report),
    "villages": ("Village / Hobli / Taluk-wise summary", village_report),
}
MONEY_REPORTS = {"compensation", "payments", "villages"}


def build_report(name, user, project_id=None):
    wb = Workbook()
    _, fn = REPORTS[name]
    fn(wb, user, _projects(user, project_id))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_farmer_statement(user, farmer):
    wb = Workbook()
    farmer_statement(wb, user, farmer)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
